"""Content ingestion (Phase G): PDF / PPTX / DOCX / TXT / manual topics.

Pipeline: input -> extraction -> cleaning -> segmentation -> concept extraction
-> (relationship extraction) -> question generation -> validation -> report.
Concept/edge proposals are provenance-stamped llm:<provider> and enter the
system as UNREVIEWED — they never steer adaptation until human approval.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .provider import GenerationSpec, LLMProvider, ProviderOutputError
from .questions import QuestionSpec, validate_batch

GRADING_BY_TYPE = {"multiple_choice": "exact_choice",
                   "numerical": "numeric_tolerance",
                   "short_answer": "normalized_match",
                   "structured": "rubric_external"}


class UnsupportedFormat(Exception):
    pass


def extract_text(payload: bytes, filename: str) -> str:
    """Format-aware text extraction. Parsers are optional dependencies; a
    missing parser raises UnsupportedFormat with an actionable message."""
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md", ""):
        return payload.decode("utf-8", errors="replace")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise UnsupportedFormat("pypdf not installed") from e
        import io
        reader = PdfReader(io.BytesIO(payload))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError as e:
            raise UnsupportedFormat("python-docx not installed") from e
        import io
        d = docx.Document(io.BytesIO(payload))
        return "\n".join(p.text for p in d.paragraphs)
    if ext == ".pptx":
        try:
            import pptx
        except ImportError as e:
            raise UnsupportedFormat("python-pptx not installed") from e
        import io
        prs = pptx.Presentation(io.BytesIO(payload))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append("\n".join(p.text for p in shape.text_frame.paragraphs))
        return "\n".join(parts)
    raise UnsupportedFormat(f"unsupported extension {ext!r}")


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def segment(text: str, max_chars: int = 1200) -> list[str]:
    """Paragraph-boundary segmentation with a hard size cap."""
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 2 > max_chars and cur:
            chunks.append(cur)
            cur = ""
        cur = (cur + "\n\n" + p).strip()
        while len(cur) > max_chars:
            chunks.append(cur[:max_chars])
            cur = cur[max_chars:]
    if cur:
        chunks.append(cur)
    return chunks


@dataclass
class IngestReport:
    provider: str
    source: str
    n_segments: int = 0
    concept_proposals: list[dict] = field(default_factory=list)
    relationship_proposals: list[dict] = field(default_factory=list)
    questions_accepted: list[QuestionSpec] = field(default_factory=list)
    questions_rejected: list[dict] = field(default_factory=list)


def ingest_text(text: str, provider: LLMProvider, source: str = "manual",
                difficulty_default: float = 0.5) -> IngestReport:
    text = clean(text)
    chunks = segment(text)
    rep = IngestReport(provider=provider.name, source=source, n_segments=len(chunks))
    seen_concepts: dict[str, dict] = {}
    for chunk in chunks:
        for prop in provider.extract_concepts(chunk):
            name = prop["name"].strip()
            if name and name not in seen_concepts:
                prop = dict(prop)
                prop["provenance"] = f"llm:{provider.name}"
                prop["human_reviewed"] = False
                seen_concepts[name] = prop
    rep.concept_proposals = list(seen_concepts.values())
    names = list(seen_concepts)
    if names:
        for rel in provider.extract_relationships(names, text):
            rel = dict(rel)
            rel["provenance"] = f"llm:{provider.name}"
            rel["human_reviewed"] = False
            rep.relationship_proposals.append(rel)
    specs = []
    for name in names:
        spec = GenerationSpec(concept=name, mastery=0.0, uncertainty=1.0,
                              difficulty=difficulty_default, intervention="EXAMPLE",
                              prerequisites=())
        try:
            q = provider.generate_question(spec)
        except ProviderOutputError as e:
            rep.questions_rejected.append({"question": f"(concept {name})",
                                           "errors": [f"provider output error: {e}"]})
            continue
        item_type = q.get("item_type", "multiple_choice")
        specs.append(QuestionSpec(
            question=q.get("question", ""), answer=str(q.get("answer", "")),
            explanation=q.get("explanation", ""), concept=name,
            difficulty=float(q.get("difficulty", difficulty_default)),
            item_type=item_type,
            prerequisites=list(q.get("prerequisites", [])),
            grading_method=q.get("grading_method") or GRADING_BY_TYPE.get(item_type, "normalized_match"),
            choices=q.get("choices"), correct_index=q.get("correct_index"),
            tolerance=q.get("tolerance"), rubric=q.get("rubric"),
            provenance=f"llm:{provider.name}"))
    ok, rejected = validate_batch(specs)
    rep.questions_accepted = ok
    rep.questions_rejected = rep.questions_rejected + rejected
    return rep
